import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]

    # Theme Palette: Warm Buddhist Saffron & Deep Navy
    COLOR_BG = RGBColor(248, 249, 250)            # Off-white background
    COLOR_PRIMARY = RGBColor(26, 32, 44)          # Deep Charcoal/Navy
    COLOR_ACCENT = RGBColor(230, 126, 34)         # Saffron / Warm Amber #E67E22
    COLOR_ACCENT_LIGHT = RGBColor(254, 243, 199)    # Saffron badge tint
    COLOR_CARD_BG = RGBColor(255, 255, 255)       # White card
    COLOR_CARD_BORDER = RGBColor(226, 232, 240)   # Slate border
    COLOR_TEXT_MAIN = RGBColor(30, 41, 59)        # Dark slate
    COLOR_TEXT_MUTED = RGBColor(100, 116, 139)    # Subdued gray
    COLOR_HEADER_BAR = RGBColor(24, 18, 14)       # Header bar dark
    COLOR_GREEN_BADGE = RGBColor(220, 252, 231)   # Soft green tint
    COLOR_GREEN_TEXT = RGBColor(22, 101, 52)      # Green text

    FONT_TITLE = "Trebuchet MS"
    FONT_BODY = "Calibri"

    def add_header(slide, title_text, slide_num_str):
        # Top banner background bar
        top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.1))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = COLOR_HEADER_BAR
        top_bar.line.color.rgb = COLOR_HEADER_BAR

        # Decorative saffron line under header
        accent_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(1.06), Inches(13.333), Inches(0.04))
        accent_line.fill.solid()
        accent_line.fill.fore_color.rgb = COLOR_ACCENT
        accent_line.line.color.rgb = COLOR_ACCENT

        # Slide Category & Title
        tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.18), Inches(9.5), Inches(0.8))
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = 0
        
        p0 = tf.paragraphs[0]
        p0.text = "BATCH-B | PROJECT REVIEW | TEAM NO. 25"
        p0.font.name = FONT_BODY
        p0.font.size = Pt(11)
        p0.font.bold = True
        p0.font.color.rgb = COLOR_ACCENT

        p1 = tf.add_paragraph()
        p1.text = title_text
        p1.font.name = FONT_TITLE
        p1.font.size = Pt(21)
        p1.font.bold = True
        p1.font.color.rgb = RGBColor(255, 255, 255)

        # Slide Number on right
        num_box = slide.shapes.add_textbox(Inches(10.5), Inches(0.25), Inches(2.0), Inches(0.6))
        ntf = num_box.text_frame
        ntf.margin_top = ntf.margin_bottom = ntf.margin_left = ntf.margin_right = 0
        np = ntf.paragraphs[0]
        np.alignment = PP_ALIGN.RIGHT
        np.text = slide_num_str
        np.font.name = FONT_TITLE
        np.font.size = Pt(24)
        np.font.bold = True
        np.font.color.rgb = COLOR_ACCENT

    def add_card(slide, left, top, width, height, title, bg_color=COLOR_CARD_BG, border_color=COLOR_CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1.5)

        if title:
            t_box = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.18), width - Inches(0.5), Inches(0.42))
            tf = t_box.text_frame
            tf.word_wrap = True
            tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = 0
            p = tf.paragraphs[0]
            p.text = title
            p.font.name = FONT_TITLE
            p.font.size = Pt(14)
            p.font.bold = True
            p.font.color.rgb = COLOR_ACCENT

        return card

    # =========================================================================
    # SLIDE 1: Title Slide
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = RGBColor(18, 14, 10)
    bg1.line.color.rgb = RGBColor(18, 14, 10)

    # Accent decorative top bar
    bar1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12))
    bar1.fill.solid()
    bar1.fill.fore_color.rgb = COLOR_ACCENT
    bar1.line.color.rgb = COLOR_ACCENT

    # Main Title Text
    tb1 = s1.shapes.add_textbox(Inches(1.0), Inches(1.1), Inches(11.333), Inches(3.2))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "BATCH-B | PROJECT REVIEW"
    p.font.name = FONT_BODY
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = COLOR_ACCENT
    p.space_after = Pt(4)

    p = tf1.add_paragraph()
    p.text = "Digitize and Showcase Monasteries of Sikkim for Tourism and Cultural Preservation"
    p.font.name = FONT_TITLE
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.space_after = Pt(10)

    p = tf1.add_paragraph()
    p.text = "A Comprehensive 360° Virtual Reality Portal, Interactive GIS Map, Festival Calendar & Cultural Heritage Platform"
    p.font.name = FONT_BODY
    p.font.size = Pt(18)
    p.font.color.rgb = RGBColor(226, 232, 240)
    p.space_after = Pt(10)

    # Team & Submission Metadata Card (Lower section)
    card_meta = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(4.5), Inches(11.333), Inches(2.3))
    card_meta.fill.solid()
    card_meta.fill.fore_color.rgb = RGBColor(28, 22, 16)
    card_meta.line.color.rgb = COLOR_ACCENT
    card_meta.line.width = Pt(1.5)

    meta_items = [
        ("TEAM NUMBER", "Team No. 25", "Batch-B"),
        ("TEAM MEMBERS", "Jakeer, Shamilini, Rushil,\nSravan, Ashritha", "5 Members (Multidisciplinary)"),
        ("DOMAIN / FOCUS", "Web & Cultural Tourism", "Monastery360 Web Portal"),
        ("SUBMISSION DATE", "26 October 2026", "Tentative Final Review")
    ]

    col_w = Inches(2.6)
    for idx, (lbl, val1, val2) in enumerate(meta_items):
        cx = Inches(1.3) + idx * (col_w + Inches(0.2))
        tbox = s1.shapes.add_textbox(cx, Inches(4.7), col_w, Inches(1.9))
        ctf = tbox.text_frame
        ctf.word_wrap = True
        
        cp0 = ctf.paragraphs[0]
        cp0.text = lbl
        cp0.font.name = FONT_BODY
        cp0.font.size = Pt(11)
        cp0.font.bold = True
        cp0.font.color.rgb = COLOR_ACCENT

        cp1 = ctf.add_paragraph()
        cp1.text = val1
        cp1.font.name = FONT_TITLE
        cp1.font.size = Pt(12.5)
        cp1.font.bold = True
        cp1.font.color.rgb = RGBColor(255, 255, 255)

        cp2 = ctf.add_paragraph()
        cp2.text = val2
        cp2.font.name = FONT_BODY
        cp2.font.size = Pt(10.5)
        cp2.font.color.rgb = RGBColor(203, 213, 225)

    # =========================================================================
    # SLIDE 2: 1. Problem Statement
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "1. Problem Statement", "02")

    # Left Box: THE PROBLEM
    add_card(s2, Inches(0.8), Inches(1.4), Inches(5.6), Inches(5.5), "THE PROBLEM")
    pb_text = s2.shapes.add_textbox(Inches(1.05), Inches(2.05), Inches(5.1), Inches(4.6))
    ptf = pb_text.text_frame
    ptf.word_wrap = True

    p_items = [
        ("Core Problem Definition",
         "• Monastery and cultural tourism information in Sikkim can be difficult to explore in one single place.\n• Information regarding locations, historical background, visiting hours, and ongoing festival events is scattered across fragmented, outdated websites."),
        ("Who Faces This Problem?",
         "• Domestic and international tourists who want to experience Sikkim's spiritual heritage but lack consolidated travel itineraries.\n• Devotees and elderly pilgrims who find remote high-altitude monasteries physically and geographically challenging to reach.\n• Cultural researchers seeking accurate archival documentation of murals, stupas, and thangkas."),
        ("Why Is Solving It Important?",
         "• Remote monasteries risk digital obscurity without modern showcase platforms.\n• Directs tourism flow sustainably to regional communities, promoting local homestays, transport providers, and rural handicrafts.")
    ]

    for q, a in p_items:
        p_q = ptf.add_paragraph() if ptf.paragraphs[0].text else ptf.paragraphs[0]
        p_q.text = q
        p_q.font.name = FONT_BODY
        p_q.font.size = Pt(12)
        p_q.font.bold = True
        p_q.font.color.rgb = COLOR_PRIMARY

        p_a = ptf.add_paragraph()
        p_a.text = a
        p_a.font.name = FONT_BODY
        p_a.font.size = Pt(11)
        p_a.font.color.rgb = COLOR_TEXT_MAIN
        p_a.space_after = Pt(12)

    # Right Box: YOUR SOLUTION & BENEFICIARIES
    add_card(s2, Inches(6.8), Inches(1.4), Inches(5.7), Inches(5.5), "YOUR SOLUTION & BENEFICIARIES")
    sol_text = s2.shapes.add_textbox(Inches(7.05), Inches(2.05), Inches(5.2), Inches(4.6))
    stf = sol_text.text_frame
    stf.word_wrap = True

    s_items = [
        ("The Monastery360 Solution",
         "• A centralized, high-performance web portal uniting 360° virtual tours, interactive geo-mapping, festival timelines, and verified accommodation directories.\n• Equirectangular panorama engine (Pannellum WebGL) bringing virtual access to sacred shrines without requiring specialized hardware."),
        ("Primary Stakeholder Benefits",
         "• Tourists & Pilgrims: Seamless exploration, panoramic 360° previews, route navigation, and integrated trip budgeting.\n• Local Communities: Increased eco-tourism revenue for regional homestays, cab operators, and authentic cultural artisans.\n• Cultural-Preservation Stakeholders: Long-term digital twin documentation of endangered architecture, centuries-old thangkas, and sacred rituals.")
    ]

    for q, a in s_items:
        p_q = stf.add_paragraph() if stf.paragraphs[0].text else stf.paragraphs[0]
        p_q.text = q
        p_q.font.name = FONT_BODY
        p_q.font.size = Pt(12)
        p_q.font.bold = True
        p_q.font.color.rgb = COLOR_ACCENT

        p_a = stf.add_paragraph()
        p_a.text = a
        p_a.font.name = FONT_BODY
        p_a.font.size = Pt(11)
        p_a.font.color.rgb = COLOR_TEXT_MAIN
        p_a.space_after = Pt(14)

    # =========================================================================
    # SLIDE 3: 2. Technology Stack & Key Features
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "2. Technology Stack & Key Features", "03")

    # Left Box: TECHNOLOGY STACK
    add_card(s3, Inches(0.8), Inches(1.4), Inches(5.6), Inches(5.5), "TECHNOLOGY STACK")
    tstack_box = s3.shapes.add_textbox(Inches(1.05), Inches(2.05), Inches(5.1), Inches(4.6))
    ttf = tstack_box.text_frame
    ttf.word_wrap = True

    tech_items = [
        ("Frontend", "HTML5, CSS3, Vanilla JavaScript (Modern ES6+ modular architecture, zero front-end framework bloat)"),
        ("Backend", "Client-side only (Ultra-lightweight, browser-based prototype ensuring instant loading and zero server maintenance cost)"),
        ("Storage", "Browser localStorage + centralized JavaScript data object (window.M360Data) for persistent user bookmarks, trip plans, and history"),
        ("Tools & Libraries", "• Pannellum (WebGL 360° equirectangular virtual tour engine)\n• Leaflet & OpenStreetMap (Interactive GIS map navigation)\n• Google Fonts (Baloo 2 & Poppins typography)\n• HTML5 Canvas (Dynamic 360° fallback placeholder generator)"),
        ("Development Tools", "VS Code, Git & GitHub repository, PowerShell optimization scripts")
    ]

    for label, desc in tech_items:
        p = ttf.add_paragraph() if ttf.paragraphs[0].text else ttf.paragraphs[0]
        p.text = f"• {label}: "
        p.font.name = FONT_BODY
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = COLOR_PRIMARY
        
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = COLOR_TEXT_MAIN
        p.space_after = Pt(8)

    # Right Box: KEY FEATURES
    add_card(s3, Inches(6.8), Inches(1.4), Inches(5.7), Inches(5.5), "KEY FEATURES")
    feat_box = s3.shapes.add_textbox(Inches(7.05), Inches(2.05), Inches(5.2), Inches(4.6))
    ftf = feat_box.text_frame
    ftf.word_wrap = True

    features = [
        ("1. 360° Virtual Tours of Monasteries", "Immersive WebGL equirectangular panoramas of 8 renowned monasteries (Rumtek, Pemayangtse, Tashiding, Enchey, Ralang, Dubdi, Dodrul, Phensang) with mouse/touch drag navigation."),
        ("2. Interactive Monastery Map", "Interactive Leaflet GIS map with GPS pin markers across East, West, South, and North Sikkim, complete with turn-by-turn Google Maps direction linking."),
        ("3. Comprehensive Festival Calendar", "Dynamic 12-month cultural calendar tracking sacred Tibetan and Sikkimese festivals (Losar, Bumchu, Saga Dawa, Pang Lhabsol) with slide-in detail drawers."),
        ("4. Cultural Art & Photo Gallery", "Curated multimedia archive showcasing golden stupas, centuries-old thangkas, and ancient vegetable-dye wall murals categorized by theme."),
        ("5. Accommodation & Travel Integration", "Integrated hotel, homestay, and travel distance listings helping visitors plan lodging near remote monasteries.")
    ]

    for title, desc in features:
        p = ftf.add_paragraph() if ftf.paragraphs[0].text else ftf.paragraphs[0]
        p.text = title
        p.font.name = FONT_BODY
        p.font.size = Pt(11.5)
        p.font.bold = True
        p.font.color.rgb = COLOR_ACCENT

        p_desc = ftf.add_paragraph()
        p_desc.text = desc
        p_desc.font.name = FONT_BODY
        p_desc.font.size = Pt(10)
        p_desc.font.color.rgb = COLOR_TEXT_MAIN
        p_desc.space_after = Pt(7)

    # =========================================================================
    # SLIDE 4: 3. Expected Outcome & Project Workflow
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "3. Project Workflow / System Architecture", "04")

    # Workflow Section (Top half)
    add_card(s4, Inches(0.8), Inches(1.35), Inches(11.733), Inches(2.7), "A. PROJECT WORKFLOW & EXPECTED OUTCOME")
    
    wf_steps = [
        ("1. INPUT", "User selects monastery from map, filters by sect/district, picks travel dates, or searches heritage keywords"),
        ("2. PROCESS", "Client-side router parses query IDs, retrieves metadata from M360Data, and prepares panorama texture assets"),
        ("3. DATA / AI", "LocalStorage loads user favorites; Pannellum WebGL decodes 360° equirectangular panorama with canvas fallbacks"),
        ("4. OUTPUT", "User immerses in 360° VR tour, views festival details, launches GPS navigation, and saves trip itinerary")
    ]

    step_w = Inches(2.6)
    for idx, (s_title, s_desc) in enumerate(wf_steps):
        sx = Inches(1.05) + idx * (step_w + Inches(0.3))
        s_card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, sx, Inches(1.95), step_w, Inches(1.85))
        s_card.fill.solid()
        s_card.fill.fore_color.rgb = COLOR_ACCENT_LIGHT if idx in (0, 3) else RGBColor(241, 245, 249)
        s_card.line.color.rgb = COLOR_ACCENT if idx in (0, 3) else COLOR_CARD_BORDER
        s_card.line.width = Pt(1.2)

        stx = s4.shapes.add_textbox(sx + Inches(0.12), Inches(2.0), step_w - Inches(0.24), Inches(1.75))
        stf = stx.text_frame
        stf.word_wrap = True
        stf.margin_top = stf.margin_bottom = stf.margin_left = stf.margin_right = 0
        
        sp = stf.paragraphs[0]
        sp.text = s_title
        sp.font.name = FONT_TITLE
        sp.font.size = Pt(12)
        sp.font.bold = True
        sp.font.color.rgb = COLOR_ACCENT if idx in (0, 3) else COLOR_PRIMARY

        sp_desc = stf.add_paragraph()
        sp_desc.text = s_desc
        sp_desc.font.name = FONT_BODY
        sp_desc.font.size = Pt(9.5)
        sp_desc.font.color.rgb = COLOR_TEXT_MAIN

        if idx < 3:
            arr_box = s4.shapes.add_textbox(sx + step_w + Inches(0.02), Inches(2.6), Inches(0.26), Inches(0.4))
            atf = arr_box.text_frame
            atf.margin_top = atf.margin_bottom = atf.margin_left = atf.margin_right = 0
            ap = atf.paragraphs[0]
            ap.text = "→"
            ap.font.size = Pt(18)
            ap.font.color.rgb = COLOR_ACCENT

    # Architecture Section (Bottom half)
    add_card(s4, Inches(0.8), Inches(4.25), Inches(11.733), Inches(2.8), "B. SIMPLE SYSTEM ARCHITECTURE")

    arch_tiers = [
        ("CLIENT / FRONTEND LAYER",
         "• 7 Responsive Pages: index.html, tours.html, stay.html, calendar.html, gallery.html, login.html, profile.html\n• Modular Components: Navbar, Search Modal (Ctrl+K), Accordion Gallery, Lightbox\n• Leaflet.js Mapping & Pannellum WebGL 360° Tour Viewer"),
        ("CLIENT-SIDE APPLICATION LOGIC",
         "• Data Router (js/app.js): Handles dynamic DOM injection and search index matching\n• Virtual Tour Engine: Equirectangular projection with auto-rotation & compass pitch\n• Dynamic Canvas Fallback: Generates synthetic panorama previews if asset is offline"),
        ("STORAGE & DATA ASSET LAYER",
         "• Centralized Heritage Data (js/data.js): 8 Monasteries, 12 Festivals, 8 Hotels\n• LocalStorage State: User session, visited gompa history, favorite bookmarks\n• Asset Pipeline: Web-optimized JPEG/PNG panoramas and thumbnails (<5.5 MB total)")
    ]

    tier_w = Inches(3.6)
    for idx, (t_title, t_desc) in enumerate(arch_tiers):
        tx = Inches(1.05) + idx * (tier_w + Inches(0.35))
        t_card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, tx, Inches(4.8), tier_w, Inches(2.05))
        t_card.fill.solid()
        t_card.fill.fore_color.rgb = RGBColor(255, 255, 255)
        t_card.line.color.rgb = COLOR_ACCENT if idx == 1 else COLOR_CARD_BORDER
        t_card.line.width = Pt(1.2)

        ttx = s4.shapes.add_textbox(tx + Inches(0.15), Inches(4.9), tier_w - Inches(0.3), Inches(1.9))
        ttf = ttx.text_frame
        ttf.word_wrap = True
        ttf.margin_top = ttf.margin_bottom = ttf.margin_left = ttf.margin_right = 0

        tp = ttf.paragraphs[0]
        tp.text = t_title
        tp.font.name = FONT_TITLE
        tp.font.size = Pt(11)
        tp.font.bold = True
        tp.font.color.rgb = COLOR_PRIMARY

        tp_desc = ttf.add_paragraph()
        tp_desc.text = t_desc
        tp_desc.font.name = FONT_BODY
        tp_desc.font.size = Pt(9.2)
        tp_desc.font.color.rgb = COLOR_TEXT_MAIN

    # =========================================================================
    # SLIDE 5: 4. Project Screenshots & Implementation Status
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "4. Project Screenshots / Output", "05")

    # Status Banner
    status_bar = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.3), Inches(11.733), Inches(0.65))
    status_bar.fill.solid()
    status_bar.fill.fore_color.rgb = RGBColor(241, 245, 249)
    status_bar.line.color.rgb = COLOR_CARD_BORDER

    st_box = s5.shapes.add_textbox(Inches(0.95), Inches(1.35), Inches(11.4), Inches(0.55))
    st_tf = st_box.text_frame
    st_tf.word_wrap = True
    st_tf.margin_top = st_tf.margin_bottom = st_tf.margin_left = st_tf.margin_right = 0
    
    stp = st_tf.paragraphs[0]
    stp.text = "IMPLEMENTATION STATUS: "
    stp.font.name = FONT_BODY
    stp.font.size = Pt(11)
    stp.font.bold = True
    stp.font.color.rgb = COLOR_PRIMARY
    
    run1 = stp.add_run()
    run1.text = "Completed: 7 core pages, 360° tours, calendar  |  Remaining: End-to-end testing, content checks & demo polish  |  Working demo ready? "
    run1.font.bold = False
    run1.font.color.rgb = COLOR_TEXT_MAIN

    run2 = stp.add_run()
    run2.text = "[In Progress / Prototype Stage]"
    run2.font.bold = True
    run2.font.color.rgb = COLOR_ACCENT

    screenshots = [
        ("SCREENSHOT 1: MAIN EXPLORER & MAP",
         "images/hero.png",
         "Interactive Monastery Map & Explorer",
         "• Real-time interactive Leaflet GIS map with 8 monastery markers.\n• Demonstrates monastery filtering, direct search, and turn-by-turn navigation."),
        ("SCREENSHOT 2: 360° VIRTUAL TOUR",
         "images/rumtek_pano.png",
         "Pannellum 360° Virtual Tour Viewer",
         "• High-resolution equirectangular panorama viewer of sacred gompas.\n• Full drag controls, pitch/yaw rotation, and fallback placeholder handling."),
        ("SCREENSHOT 3: CULTURAL CALENDAR & STAYS",
         "images/pemayangtse.jpg",
         "Festival Calendar & Stay Filters",
         "• Demonstrates interactive 12-month Buddhist festival timeline.\n• Showcases cultural gallery and verified nearby hotel/homestay directory.")
    ]

    card_w = Inches(3.65)
    for idx, (stitle, img_rel_path, caption, desc) in enumerate(screenshots):
        cx = Inches(0.8) + idx * (card_w + Inches(0.39))
        card = add_card(s5, cx, Inches(2.1), card_w, Inches(5.0), stitle)

        img_path = os.path.abspath(img_rel_path)
        if os.path.exists(img_path):
            try:
                s5.shapes.add_picture(img_path, cx + Inches(0.2), Inches(2.65), card_w - Inches(0.4), Inches(2.0))
            except Exception as e:
                print(f"Image load warning: {e}")
        
        cap_box = s5.shapes.add_textbox(cx + Inches(0.2), Inches(4.8), card_w - Inches(0.4), Inches(2.1))
        ctf = cap_box.text_frame
        ctf.word_wrap = True
        ctf.margin_top = ctf.margin_bottom = ctf.margin_left = ctf.margin_right = 0

        cp0 = ctf.paragraphs[0]
        cp0.text = caption
        cp0.font.name = FONT_TITLE
        cp0.font.size = Pt(11.5)
        cp0.font.bold = True
        cp0.font.color.rgb = COLOR_PRIMARY
        cp0.space_after = Pt(3)

        cp1 = ctf.add_paragraph()
        cp1.text = desc
        cp1.font.name = FONT_BODY
        cp1.font.size = Pt(9.5)
        cp1.font.color.rgb = COLOR_TEXT_MAIN

    # =========================================================================
    # SLIDE 6: 5. Team Members & Contributions
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "5. Team Members & Contributions", "06")

    add_card(s6, Inches(0.8), Inches(1.35), Inches(11.733), Inches(5.6), "TEAM NO. 25 CONTRIBUTIONS MATRIX")

    rows = 6
    cols = 3
    left = Inches(1.1)
    top = Inches(2.0)
    width = Inches(11.133)
    height = Inches(4.7)

    table_shape = s6.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table

    table.columns[0].width = Inches(2.3)
    table.columns[1].width = Inches(3.4)
    table.columns[2].width = Inches(5.433)

    headers = ["TEAM MEMBER", "TASK / MODULE", "WHAT THEY CONTRIBUTED"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_PRIMARY
        for p in cell.text_frame.paragraphs:
            p.alignment = PP_ALIGN.LEFT
            p.font.name = FONT_TITLE
            p.font.size = Pt(11.5)
            p.font.bold = True
            p.font.color.rgb = RGBColor(255, 255, 255)

    team_data = [
        ("Jakeer",
         "Accommodation & Travel Management",
         "• Developed stay.html directory and nearby accommodation filters (Luxury, Heritage, Boutique, Budget).\n• Integrated distance calculations and transport guide for travelers near monasteries."),
        ("Shamilini",
         "Virtual Tour Development",
         "• Implemented Pannellum.js WebGL 360° virtual tour viewer across 8 monasteries.\n• Handled equirectangular panorama image loading, thumbnail switching, and fallback preview canvas."),
        ("Rushil",
         "Authentication & UI Integration",
         "• Engineered dual-mode authentication (localStorage session state + optional Supabase cloud schema).\n• Integrated responsive navigation bar, search modal (Ctrl+K), and core multi-page UI architecture."),
        ("Sravan",
         "Interactive Map & Navigation",
         "• Built Leaflet.js interactive GIS map with custom monastery markers across Sikkim districts.\n• Implemented coordinate popups with one-click direct Google Maps turn-by-turn navigation."),
        ("Ashritha",
         "Festival Calendar & Cultural Gallery",
         "• Designed calendar.html interactive monthly timeline for sacred events (Losar, Bumchu, Saga Dawa).\n• Created cultural gallery (gallery.html) with category filters (Art, History, Architecture) and lightbox.")
    ]

    for row_idx, data in enumerate(team_data, start=1):
        for col_idx, text in enumerate(data):
            cell = table.cell(row_idx, col_idx)
            cell.text = text
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(255, 255, 255) if row_idx % 2 == 1 else RGBColor(248, 250, 252)
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.LEFT
                p.font.name = FONT_BODY
                p.font.size = Pt(9.8)
                if col_idx == 0:
                    p.font.bold = True
                    p.font.color.rgb = COLOR_PRIMARY
                elif col_idx == 1:
                    p.font.bold = True
                    p.font.color.rgb = COLOR_ACCENT
                else:
                    p.font.color.rgb = COLOR_TEXT_MAIN

    # =========================================================================
    # SLIDE 7: 6. AI Project Video & Review Talking Points
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "6. AI Project Video & Conclusion", "07")

    # Left Box: MANDATORY AI VIDEO SUBMISSION
    add_card(s7, Inches(0.8), Inches(1.35), Inches(5.6), Inches(5.6), "MANDATORY AI VIDEO SUBMISSION")
    vbox = s7.shapes.add_textbox(Inches(1.05), Inches(1.95), Inches(5.1), Inches(4.8))
    vtf = vbox.text_frame
    vtf.word_wrap = True

    vp0 = vtf.paragraphs[0]
    vp0.text = "VIDEO DURATION: 1 MIN 20 SEC  |  AI VOICE-OVER"
    vp0.font.name = FONT_TITLE
    vp0.font.size = Pt(12)
    vp0.font.bold = True
    vp0.font.color.rgb = COLOR_ACCENT
    vp0.space_after = Pt(6)

    v_timeline = [
        ("0:00 – 0:20 | Problem & Heritage Context",
         "Introduces Sikkim's remote monasteries and the challenge of accessing scattered cultural information."),
        ("0:20 – 0:45 | Virtual Tour & Map Demo",
         "Live demonstration of selecting a monastery, launching the 360° Pannellum viewer, and exploring the GIS map."),
        ("0:45 – 1:05 | Festival Calendar & Gallery",
         "Showcasing annual ritual events (Losar, Bumchu) and filtering the sacred art and thangka gallery."),
        ("1:05 – 1:20 | Benefits & Conclusion",
         "Summary of cultural preservation benefits and impact on sustainable regional tourism.")
    ]

    for timing, desc in v_timeline:
        p_t = vtf.add_paragraph()
        p_t.text = timing
        p_t.font.name = FONT_BODY
        p_t.font.size = Pt(10.5)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_PRIMARY

        p_d = vtf.add_paragraph()
        p_d.text = desc
        p_d.font.name = FONT_BODY
        p_d.font.size = Pt(9.8)
        p_d.font.color.rgb = COLOR_TEXT_MAIN
        p_d.space_after = Pt(4)

    # Video link placeholder box
    vlink_shape = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.05), Inches(5.8), Inches(5.1), Inches(0.95))
    vlink_shape.fill.solid()
    vlink_shape.fill.fore_color.rgb = COLOR_ACCENT_LIGHT
    vlink_shape.line.color.rgb = COLOR_ACCENT
    vlink_shape.line.width = Pt(1.2)

    vltx = s7.shapes.add_textbox(Inches(1.15), Inches(5.85), Inches(4.9), Inches(0.85))
    vltf = vltx.text_frame
    vltf.word_wrap = True
    vltf.margin_top = vltf.margin_bottom = vltf.margin_left = vltf.margin_right = 0
    vlp0 = vltf.paragraphs[0]
    vlp0.text = "AI Video Submission Link / File Name:"
    vlp0.font.name = FONT_BODY
    vlp0.font.size = Pt(10)
    vlp0.font.bold = True
    vlp0.font.color.rgb = COLOR_PRIMARY

    vlp1 = vltf.add_paragraph()
    vlp1.text = "Monastery360_AI_Voiceover_Demo_1m20s.mp4"
    vlp1.font.name = FONT_TITLE
    vlp1.font.size = Pt(10.5)
    vlp1.font.color.rgb = COLOR_ACCENT

    # Right Box: REVIEW TALKING POINTS & CONCLUSION
    add_card(s7, Inches(6.8), Inches(1.35), Inches(5.7), Inches(5.6), "REVIEW TALKING POINTS & CONCLUSION")
    cbox = s7.shapes.add_textbox(Inches(7.05), Inches(1.95), Inches(5.2), Inches(4.8))
    ctf = cbox.text_frame
    ctf.word_wrap = True

    talking_points = [
        ("Why This Stack?",
         "Lightweight, browser-based prototype requiring zero server overhead, loading instantly on any standard device."),
        ("What Is Working?",
         "360° virtual tour viewer, interactive Leaflet map, cultural gallery, festival calendar, and tourism stay pages."),
        ("Key Challenge Identified",
         "Creating awareness about lesser-known remote monasteries and collecting authentic high-resolution 360° assets."),
        ("Remaining Work Ahead",
         "Adding more monasteries to the interactive map and capturing more panoramic street views for full coverage."),
        ("Tentative Final Submission Date",
         "26 October 2026")
    ]

    for q, a in talking_points:
        p_q = ctf.add_paragraph() if ctf.paragraphs[0].text else ctf.paragraphs[0]
        p_q.text = f"• {q}: "
        p_q.font.name = FONT_BODY
        p_q.font.size = Pt(10.5)
        p_q.font.bold = True
        p_q.font.color.rgb = COLOR_ACCENT if q.startswith("Tentative") else COLOR_PRIMARY

        run = p_q.add_run()
        run.text = a
        run.font.bold = True if q.startswith("Tentative") else False
        run.font.color.rgb = COLOR_ACCENT if q.startswith("Tentative") else COLOR_TEXT_MAIN
        p_q.space_after = Pt(5)

    # Thank you footer bar inside right box
    ty_box = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.05), Inches(5.8), Inches(5.2), Inches(0.95))
    ty_box.fill.solid()
    ty_box.fill.fore_color.rgb = COLOR_PRIMARY
    ty_box.line.color.rgb = COLOR_PRIMARY
    
    tytx = s7.shapes.add_textbox(Inches(7.15), Inches(5.95), Inches(5.0), Inches(0.65))
    tytf = tytx.text_frame
    tytf.margin_top = tytf.margin_bottom = tytf.margin_left = tytf.margin_right = 0
    typ = tytf.paragraphs[0]
    typ.alignment = PP_ALIGN.CENTER
    typ.text = "Thank You  |  Questions & Answers"
    typ.font.name = FONT_TITLE
    typ.font.size = Pt(16)
    typ.font.bold = True
    typ.font.color.rgb = RGBColor(255, 255, 255)

    # Save presentation
    target_path = os.path.abspath("Monastery360_Final_Project_Review.pptx")
    prs.save(target_path)
    print(f"SUCCESS: Presentation saved at {target_path}")

    # Also try updating Monastery360_Project_Presentation.pptx if unlocked
    alt_path = os.path.abspath("Monastery360_Project_Presentation.pptx")
    try:
        prs.save(alt_path)
        print(f"SUCCESS: Also updated {alt_path}")
    except Exception as e:
        print(f"Notice: {alt_path} is currently open in PowerPoint ({e}). Saved cleanly to {target_path}")

if __name__ == "__main__":
    create_deck()
